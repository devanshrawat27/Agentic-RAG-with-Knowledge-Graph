"""Document loader: read PDF / DOCX / PPTX / TXT bytes into plain text."""

from io import BytesIO
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt", ".md"}


class UnsupportedDocumentError(ValueError):
    pass


def _extension(filename: str) -> str:
    return Path(filename).suffix.lower()


def load_txt(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def load_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages)


def load_docx(data: bytes) -> str:
    import docx

    document = docx.Document(BytesIO(data))
    return "\n".join(p.text for p in document.paragraphs)


def load_pptx(data: bytes) -> str:
    from pptx import Presentation

    presentation = Presentation(BytesIO(data))
    lines: list[str] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    lines.append(text)
    return "\n".join(lines)


def load_document(filename: str, data: bytes) -> str:
    """Extract text from a document's raw bytes based on its extension."""
    ext = _extension(filename)
    if ext == ".pdf":
        text = load_pdf(data)
    elif ext == ".docx":
        text = load_docx(data)
    elif ext == ".pptx":
        text = load_pptx(data)
    elif ext in {".txt", ".md"}:
        text = load_txt(data)
    else:
        raise UnsupportedDocumentError(
            f"Unsupported file type '{ext}'. Supported: {sorted(SUPPORTED_EXTENSIONS)}"
        )
    return _clean(text)


def load_path(path: str | Path) -> str:
    path = Path(path)
    return load_document(path.name, path.read_bytes())


def _clean(text: str) -> str:
    lines = [line.rstrip() for line in text.replace("\r\n", "\n").split("\n")]
    collapsed: list[str] = []
    blank = False
    for line in lines:
        if line.strip() == "":
            if not blank:
                collapsed.append("")
            blank = True
        else:
            collapsed.append(line)
            blank = False
    return "\n".join(collapsed).strip()
